#!/usr/bin/env python3
"""
Generate a conference abstract DOCX for the pLIN system.
Output: output/pLIN_Conference_Abstract.docx
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

OUTPUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "output")
os.makedirs(OUTPUT_DIR, exist_ok=True)


def set_cell_shading(cell, color_hex):
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(
        qn("w:shd"),
        {qn("w:fill"): color_hex, qn("w:val"): "clear"},
    )
    shading.append(shading_elem)


def add_run(para, text, bold=False, italic=False, size=11, font="Times New Roman", color=None, superscript=False):
    run = para.add_run(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.superscript = superscript
    if color:
        run.font.color.rgb = RGBColor(*color)
    return run


def main():
    doc = Document()

    # Page setup
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ── Title ────────────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    add_run(title,
            "pLIN: a hierarchical plasmid classification and surveillance system "
            "linking antimicrobial resistance to transmissible lineages",
            bold=True, size=12)

    # ── Authors ──────────────────────────────────────────────────────────────
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(2)
    add_run(authors, "Basil Britto Xavier", bold=False, size=10, superscript=False)
    add_run(authors, "1", size=8, superscript=True)
    add_run(authors, ", Anurag Kumar Bari", size=10)
    add_run(authors, "1", size=8, superscript=True)
    add_run(authors, ", Bhanu Sinha", size=10)
    add_run(authors, "1", size=8, superscript=True)
    add_run(authors, ", John W.A. Rossen", size=10)
    add_run(authors, "1", size=8, superscript=True)
    add_run(authors, " on behalf of the DRAIGON Consortium", size=10, italic=True)

    # ── Affiliations ─────────────────────────────────────────────────────────
    affil = doc.add_paragraph()
    affil.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affil.paragraph_format.space_after = Pt(12)
    add_run(affil, "1", size=8, superscript=True)
    add_run(affil,
            " Department of Medical Microbiology and Infection Prevention, "
            "University Medical Center Groningen, University of Groningen, "
            "Groningen, The Netherlands",
            size=9, italic=True)

    # ── Separator ────────────────────────────────────────────────────────────
    sep = doc.add_paragraph()
    sep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sep.paragraph_format.space_before = Pt(0)
    sep.paragraph_format.space_after = Pt(8)

    # ── ABSTRACT BODY (structured) ───────────────────────────────────────────

    # Background
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "Background: ", bold=True, size=11)
    add_run(p,
            "Plasmid-mediated horizontal gene transfer drives antimicrobial resistance (AMR) "
            "dissemination in clinically important bacteria. Current typing approaches "
            "(PlasmidFinder, pMLST, MOB-suite, COPLA, mge-cluster) lack the resolution, "
            "permanence, and integrated AMR profiling needed to track individual "
            "resistance-carrying plasmid lineages across healthcare settings and national borders.",
            size=11)

    # Methods
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "Methods: ", bold=True, size=11)
    add_run(p,
            "We developed pLIN (plasmid Lineage Identification Number), a hierarchical "
            "classification system assigning permanent, six-level codes to plasmids based on "
            "tetranucleotide (4-mer) frequency vectors (256 features) and single-linkage "
            "clustering at six cosine distance thresholds calibrated against average nucleotide "
            "identity (ANI; L1 ~85% to L6 ~99.9%). A k-nearest-neighbour classifier (k=5, "
            "cosine metric, distance-weighted) was trained on 8,077 complete plasmid sequences "
            "across 28 Inc/Rep groups: 20 Gram-negative, 4 Gram-positive (", size=11)
    add_run(p, "Staphylococcus aureus, Enterococcus", italic=True, size=11)
    add_run(p, "), 2 ", size=11)
    add_run(p, "Acinetobacter baumannii", italic=True, size=11)
    add_run(p, ", and 2 ", size=11)
    add_run(p, "Pseudomonas aeruginosa", italic=True, size=11)
    add_run(p,
            " rep type groups. AMR profiling was performed with AMRFinderPlus v4.2.5. "
            "Seven analytical modules address assembly completeness, database coverage, "
            "recombination detection, novel group discovery, evolutionary rate estimation, "
            "cluster stability, and mobile genetic element (MGE) boundary detection. "
            "Validation used FastANI across 4,970 plasmid pairs and cross-validation against "
            "74 plasmids from 27 published outbreak studies spanning 13 countries.",
            size=11)

    # Results
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "Results: ", bold=True, size=11)
    add_run(p,
            "pLIN resolved 3,073 unique strain-level codes from 8,077 plasmids, achieving a "
            "Simpson\u2019s diversity index of 0.985 versus 0.641 for Inc typing alone "
            "(1.54-fold improvement). The KNN classifier achieved 91.1% accuracy across "
            "28 groups (five-fold cross-validation). AMRFinderPlus detected 64,891 gene hits "
            "across 83.1% of Gram-negative plasmids, including 1,635 carbapenemase, "
            "1,804 ESBL, 204 mobile colistin resistance (", size=11)
    add_run(p, "mcr", italic=True, size=11)
    add_run(p,
            "), and 2,315 plasmid-mediated quinolone resistance determinants. "
            "pLIN identified high-risk lineages invisible to conventional typing: "
            "pLIN 671 (IncN; n=90; 100% ", size=11)
    add_run(p, "bla", italic=True, size=11)
    add_run(p, "KPC-2", size=9, superscript=False)
    add_run(p, " carriage; mean 13.2 AMR genes) and pLIN 860 (five Inc groups; n=142; "
            "mean 14.4 AMR genes; 44.4% ", size=11)
    add_run(p, "mcr", italic=True, size=11)
    add_run(p,
            " carriage). Cross-validation against 27 outbreak studies confirmed correct "
            "identification of known high-risk lineages with 85.1% high-confidence "
            "classifications. Expansion to 79,305 plasmids yielded 57,886 unique codes "
            "with a 97.3% classification rate.",
            size=11)

    # Conclusions
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_run(p, "Conclusions: ", bold=True, size=11)
    add_run(p,
            "pLIN provides the first permanent, hierarchical classification system for "
            "plasmids that integrates multi-resolution typing with AMR surveillance and "
            "automated outbreak detection across Gram-negative, Gram-positive, and WHO "
            "critical priority pathogens. The open-source tool runs on standard hardware "
            "in under 30 minutes, requiring no specialist bioinformatics expertise, and is "
            "immediately deployable for routine plasmid surveillance in clinical microbiology "
            "laboratories. pLIN enables prospective tracking of high-risk plasmid lineages "
            "across institutions, supporting integration into national and international "
            "surveillance networks.",
            size=11)

    # ── Keywords ─────────────────────────────────────────────────────────────
    kw = doc.add_paragraph()
    kw.paragraph_format.space_before = Pt(8)
    kw.paragraph_format.space_after = Pt(4)
    add_run(kw, "Keywords: ", bold=True, size=10)
    add_run(kw,
            "plasmid classification, antimicrobial resistance, hierarchical clustering, "
            "genomic epidemiology, surveillance, tetranucleotide composition",
            size=10, italic=True)

    # ── Word count note ──────────────────────────────────────────────────────
    wc = doc.add_paragraph()
    wc.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    wc.paragraph_format.space_before = Pt(12)

    # Calculate approximate word count of abstract body
    body_text = (
        "Background: Plasmid-mediated horizontal gene transfer drives antimicrobial resistance (AMR) "
        "dissemination in clinically important bacteria. Current typing approaches "
        "(PlasmidFinder, pMLST, MOB-suite, COPLA, mge-cluster) lack the resolution, "
        "permanence, and integrated AMR profiling needed to track individual "
        "resistance-carrying plasmid lineages across healthcare settings and national borders. "
        "Methods: We developed pLIN (plasmid Lineage Identification Number), a hierarchical "
        "classification system assigning permanent, six-level codes to plasmids based on "
        "tetranucleotide (4-mer) frequency vectors (256 features) and single-linkage "
        "clustering at six cosine distance thresholds calibrated against average nucleotide "
        "identity (ANI; L1 ~85% to L6 ~99.9%). A k-nearest-neighbour classifier (k=5, "
        "cosine metric, distance-weighted) was trained on 8,077 complete plasmid sequences "
        "across 28 Inc/Rep groups: 20 Gram-negative, 4 Gram-positive (Staphylococcus aureus, "
        "Enterococcus), 2 Acinetobacter baumannii, and 2 Pseudomonas aeruginosa "
        "rep type groups. AMR profiling was performed with AMRFinderPlus v4.2.5. "
        "Seven analytical modules address assembly completeness, database coverage, "
        "recombination detection, novel group discovery, evolutionary rate estimation, "
        "cluster stability, and mobile genetic element (MGE) boundary detection. "
        "Validation used FastANI across 4,970 plasmid pairs and cross-validation against "
        "74 plasmids from 27 published outbreak studies spanning 13 countries. "
        "Results: pLIN resolved 3,073 unique strain-level codes from 8,077 plasmids, achieving a "
        "Simpson's diversity index of 0.985 versus 0.641 for Inc typing alone "
        "(1.54-fold improvement). The KNN classifier achieved 91.1% accuracy across "
        "28 groups (five-fold cross-validation). AMRFinderPlus detected 64,891 gene hits "
        "across 83.1% of Gram-negative plasmids, including 1,635 carbapenemase, "
        "1,804 ESBL, 204 mobile colistin resistance (mcr), and 2,315 plasmid-mediated "
        "quinolone resistance determinants. pLIN identified high-risk lineages invisible "
        "to conventional typing: pLIN 671 (IncN; n=90; 100% blaKPC-2 carriage; mean 13.2 "
        "AMR genes) and pLIN 860 (five Inc groups; n=142; mean 14.4 AMR genes; 44.4% mcr "
        "carriage). Cross-validation against 27 outbreak studies confirmed correct "
        "identification of known high-risk lineages with 85.1% high-confidence "
        "classifications. Expansion to 79,305 plasmids yielded 57,886 unique codes "
        "with a 97.3% classification rate. "
        "Conclusions: pLIN provides the first permanent, hierarchical classification system for "
        "plasmids that integrates multi-resolution typing with AMR surveillance and "
        "automated outbreak detection across Gram-negative, Gram-positive, and WHO "
        "critical priority pathogens. The open-source tool runs on standard hardware "
        "in under 30 minutes, requiring no specialist bioinformatics expertise, and is "
        "immediately deployable for routine plasmid surveillance in clinical microbiology "
        "laboratories. pLIN enables prospective tracking of high-risk plasmid lineages "
        "across institutions, supporting integration into national and international "
        "surveillance networks."
    )
    word_count = len(body_text.split())
    add_run(wc, f"Word count: {word_count}", size=9, italic=True, color=(128, 128, 128))

    # Save
    out_path = os.path.join(OUTPUT_DIR, "pLIN_Conference_Abstract.docx")
    doc.save(out_path)
    print(f"Conference abstract saved to: {out_path}")
    print(f"Abstract word count: {word_count}")


if __name__ == "__main__":
    main()
